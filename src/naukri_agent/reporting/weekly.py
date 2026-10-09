"""
The weekly report: one Excel file with every job the app gave you during a Monday-to-Sunday week and where each one
stands. Built every Sunday at 22:00 by `naukri-agent weekly-report` and emailed to you as an attachment (a copy is kept
on the PC under settings.weekly_report_dir).

"Given" means:
  * Company website: a job that was in a digest's Part 1 that week.
  * Naukri Apply via Telegram: a job the app offered you through `telegram-apply` that week, or that a Telegram ping
    listed that week. (The ping's job list is recorded from 2026-10-09; before that only the jobs actually offered
    are known.)

Status is as of the moment the report is made. It reads the database only: no browser, no Naukri, nothing is applied or
changed, and nothing here is invented: a status is what the records say.
"""

from __future__ import annotations

import datetime
import logging
from dataclasses import dataclass, field
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from naukri_agent.config import Settings

logger = logging.getLogger(__name__)

ROUTE_COMPANY = "Company website"
ROUTE_TELEGRAM = "Naukri Apply (Telegram)"

_JOBS_HEADERS = [
    "Date given", "Day", "Route", "Job title", "Company", "Location", "Match score", "Suggested resume",
    "Status", "Status date", "Direct apply link", "Naukri link", "Notes",
]
_GREEN, _YELLOW, _GREY, _RED = "C6EFCE", "FFEB9C", "E7E6E6", "F4CCCC"

_ATTEMPT_STATUS = {  # AutoApplyAttempt.outcome -> what it means for you
    "declined": ("you said no", _GREY),
    "no_reply": ("no reply on Telegram - not applied", _GREY),
    "failed": ("failed - not applied", _RED),
    "unconfirmed": ("unconfirmed - check on Naukri", _RED),
    "needs_human": ("needed your input - not applied", _YELLOW),
    "not_native": ("no Naukri Apply button any more", _GREY),
    "excluded_type": ("contract or temporary job - skipped", _GREY),
}


@dataclass
class WeeklyRow:
    given_on: datetime.date
    route: str
    title: str
    company: str
    location: str | None
    score: float | None
    resume: str | None
    status: str
    colour: str
    status_date: datetime.date | None = None
    direct_link: str | None = None
    naukri_url: str = ""
    note: str = ""


@dataclass
class WeeklyReport:
    week_start: datetime.date  # the Monday
    week_end: datetime.date  # the Sunday
    generated_local: datetime.datetime
    rows: list[WeeklyRow] = field(default_factory=list)
    path: Path | None = None

    def counts_by_route(self) -> dict[str, int]:
        out: dict[str, int] = {}
        for r in self.rows:
            out[r.route] = out.get(r.route, 0) + 1
        return out

    def counts_by_status(self) -> dict[str, int]:
        out: dict[str, int] = {}
        for r in self.rows:
            out[r.status] = out.get(r.status, 0) + 1
        return dict(sorted(out.items(), key=lambda kv: (-kv[1], kv[0])))


# --- which week ---------------------------------------------------------------------------------------------------------


def week_bounds(today: datetime.date, *, week_of: datetime.date | None = None) -> tuple[datetime.date, datetime.date]:
    """(Monday, Sunday). `week_of` picks the week containing that date. Otherwise: on a Sunday, this week; on any other day
    (a missed Sunday caught up later), the most recent Monday-to-Sunday week that has finished."""
    if week_of is not None:
        monday = week_of - datetime.timedelta(days=week_of.weekday())
    elif today.weekday() == 6:
        monday = today - datetime.timedelta(days=6)
    else:
        monday = today - datetime.timedelta(days=today.weekday() + 7)
    return monday, monday + datetime.timedelta(days=6)


def _utc_naive(local_date: datetime.date, tz: ZoneInfo) -> datetime.datetime:
    midnight = datetime.datetime(local_date.year, local_date.month, local_date.day, tzinfo=tz)
    return midnight.astimezone(datetime.UTC).replace(tzinfo=None)


def _local_date(moment: datetime.datetime | None, tz: ZoneInfo) -> datetime.date | None:
    if moment is None:
        return None
    return moment.replace(tzinfo=datetime.UTC).astimezone(tz).date()


# --- collecting -----------------------------------------------------------------------------------------------------------


def _company_status(session: Session, job_id: int) -> tuple[str, str, datetime.datetime | None]:
    from naukri_agent.database.models import ApplicationHistory, ApplicationStatus, FollowupPrompt
    from naukri_agent.database.repositories import application_label

    row = session.query(ApplicationHistory).filter_by(job_id=job_id).one_or_none()
    status = row.status if row is not None else ApplicationStatus.NOT_APPLIED
    if status in (ApplicationStatus.NOT_APPLIED, ApplicationStatus.UNKNOWN):
        prompt = session.query(FollowupPrompt).filter_by(job_id=job_id).one_or_none()
        later = f" (put off {prompt.later_count}x)" if prompt is not None and prompt.later_count else ""
        return "waiting for your answer" + later, _YELLOW, None
    when = row.applied_at or row.updated_at
    label = application_label(row)
    colour = _GREEN if status == ApplicationStatus.APPLIED else _GREY if status in (ApplicationStatus.NOT_APPLYING, ApplicationStatus.IGNORED) else _GREEN
    return label, colour, when


def _telegram_status(session: Session, job_id: int, week_attempts: list) -> tuple[str, str, datetime.datetime | None]:
    from naukri_agent.database.models import ApplicationHistory, ApplicationStatus
    from naukri_agent.database.repositories import application_label

    row = session.query(ApplicationHistory).filter_by(job_id=job_id).one_or_none()
    if row is not None and row.status not in (ApplicationStatus.NOT_APPLIED, ApplicationStatus.UNKNOWN):
        colour = _GREY if row.status in (ApplicationStatus.NOT_APPLYING, ApplicationStatus.IGNORED) else _GREEN
        return application_label(row), colour, row.applied_at or row.updated_at
    if week_attempts:
        last = week_attempts[-1]
        text, colour = _ATTEMPT_STATUS.get(last.outcome, (last.outcome, _YELLOW))
        return text, colour, last.attempted_at
    return "ready on Telegram - not offered yet", _YELLOW, None


def collect_week(session: Session, settings: Settings, monday: datetime.date, sunday: datetime.date) -> list[WeeklyRow]:
    from naukri_agent.database.models import (
        AutoApplyAttempt,
        DailyRun,
        Job,
        JobMatch,
        JobRecommendation,
        ResumeSelection,
        RunEvent,
    )
    from naukri_agent.database.repositories import canonical_job_id

    tz = ZoneInfo(settings.timezone)
    start, end = _utc_naive(monday, tz), _utc_naive(sunday + datetime.timedelta(days=1), tz)

    # Everything that "gave" you a job this week, keyed by the canonical job id, with the earliest moment it was given.
    given: dict[int, datetime.datetime] = {}
    rec_of: dict[int, JobRecommendation] = {}
    for rec in (
        session.query(JobRecommendation)
        .filter(JobRecommendation.recommended_at >= start, JobRecommendation.recommended_at < end)
        .order_by(JobRecommendation.recommended_at.asc())
    ):
        canon = canonical_job_id(session, rec.job_id)
        given.setdefault(canon, rec.recommended_at)
        rec_of.setdefault(canon, rec)
    for ev, run in (
        session.query(RunEvent, DailyRun)
        .join(DailyRun, DailyRun.id == RunEvent.daily_run_id)
        .filter(RunEvent.stage == "telegram_ping", DailyRun.started_at >= start, DailyRun.started_at < end)
        .order_by(DailyRun.started_at.asc())
    ):
        for jid in (ev.detail or {}).get("job_ids") or []:
            given.setdefault(canonical_job_id(session, int(jid)), run.started_at)
    attempts: dict[int, list] = {}
    for att in (
        session.query(AutoApplyAttempt)
        .filter(AutoApplyAttempt.attempted_at >= start, AutoApplyAttempt.attempted_at < end)
        .order_by(AutoApplyAttempt.attempted_at.asc())
    ):
        canon = canonical_job_id(session, att.job_id)
        attempts.setdefault(canon, []).append(att)
        given.setdefault(canon, att.attempted_at)

    rows: list[WeeklyRow] = []
    for canon, first in given.items():
        job = session.get(Job, canon)
        if job is None:
            continue
        match = session.query(JobMatch).filter_by(job_id=canon).order_by(JobMatch.matched_at.desc()).first()
        sel = session.query(ResumeSelection).filter_by(job_id=canon).order_by(ResumeSelection.selected_at.desc()).first()
        rec = rec_of.get(canon)
        score = rec.score_at_email if rec is not None else (match.overall_score if match else None)
        resume = (rec.resume_id_at_email if rec is not None else None) or (sel.resume_id if sel else None)
        # The route is the job's real type, not which email it happened to be in (older digests listed every job).
        if job.apply_type == "native":
            route = ROUTE_TELEGRAM
            status, colour, when = _telegram_status(session, canon, attempts.get(canon, []))
            direct = None
        else:
            route = ROUTE_COMPANY
            status, colour, when = _company_status(session, canon)
            direct = job.apply_redirect_url
        note = ""
        if colour == _RED:
            note = "Check this job on Naukri: the app could not confirm the result, and it may still have gone through."
        rows.append(WeeklyRow(
            given_on=_local_date(first, tz), route=route, title=job.title, company=job.company, location=job.location,
            score=score, resume=resume, status=status, colour=colour, status_date=_local_date(when, tz),
            direct_link=direct, naukri_url=job.url, note=note,
        ))
    return sorted(rows, key=lambda r: (r.given_on, r.route, r.company))


# --- the workbook ---------------------------------------------------------------------------------------------------------


def _summary_sheet(ws, report: WeeklyReport) -> None:
    bold = Font(bold=True)
    ws.append([f"Weekly job report: {report.week_start:%a %d %b} to {report.week_end:%a %d %b %Y}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Made {report.generated_local:%a %d %b %Y, %I:%M %p}. Statuses are as of that moment."])
    ws.append([])
    ws.append(["Jobs given this week", len(report.rows)])
    ws.cell(row=ws.max_row, column=1).font = bold
    ws.append([])
    ws.append(["By route", "Jobs"])
    for c in ws[ws.max_row]:
        c.font = bold
    for route, n in report.counts_by_route().items():
        ws.append([route, n])
    ws.append([])
    ws.append(["By status", "Jobs"])
    for c in ws[ws.max_row]:
        c.font = bold
    for status, n in report.counts_by_status().items():
        ws.append([status, n])
    ws.append([])
    ws.append(["Company website = in Part 1 of a digest. Naukri Apply (Telegram) = offered by telegram-apply or listed in a Telegram ping."])
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 12


def _jobs_sheet(ws, report: WeeklyReport) -> None:
    ws.append(_JOBS_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
    for r in report.rows:
        ws.append([
            r.given_on, r.given_on.strftime("%a"), r.route, r.title, r.company, r.location or "",
            round(r.score, 1) if r.score is not None else "", r.resume or "", r.status,
            r.status_date or "", r.direct_link or "", r.naukri_url, r.note,
        ])
        row = ws.max_row
        ws.cell(row=row, column=1).number_format = "dd mmm yyyy"
        ws.cell(row=row, column=10).number_format = "dd mmm yyyy"
        ws.cell(row=row, column=9).fill = PatternFill("solid", fgColor=r.colour)
        for col, url in ((11, r.direct_link), (12, r.naukri_url)):
            if url:
                ws.cell(row=row, column=col).hyperlink = url
                ws.cell(row=row, column=col).font = Font(color="0563C1", underline="single")
    widths = [13, 6, 24, 38, 28, 24, 12, 18, 38, 13, 46, 46, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top")


def write_workbook(report: WeeklyReport, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _summary_sheet(wb.active, report)
    wb.active.title = "Summary"
    _jobs_sheet(wb.create_sheet("Jobs"), report)
    wb.save(path)
    report.path = path
    return path


def report_filename(monday: datetime.date, sunday: datetime.date) -> str:
    return f"weekly_{monday:%Y-%m-%d}_to_{sunday:%Y-%m-%d}.xlsx"


def email_text(report: WeeklyReport) -> tuple[str, str]:
    n = len(report.rows)
    subject = f"[naukri-agent] Weekly report: {report.week_start:%d %b} to {report.week_end:%d %b %Y} - {n} job{'s' if n != 1 else ''}"
    lines = [f"Weekly job report: {report.week_start:%a %d %b} to {report.week_end:%a %d %b %Y}", ""]
    if not n:
        lines.append("No jobs were given this week.")
    else:
        lines.append(f"{n} job{'s were' if n != 1 else ' was'} given this week.")
        lines += ["", "By route:"] + [f"  {route}: {k}" for route, k in report.counts_by_route().items()]
        lines += ["", "By status:"] + [f"  {status}: {k}" for status, k in report.counts_by_status().items()]
    lines += ["", f"Every job, with its link and status, is in the attached Excel file ({report.path.name if report.path else ''}).",
              f"Statuses are as of {report.generated_local:%a %d %b, %I:%M %p}."]
    return subject, "\n".join(lines)


# --- the whole job -------------------------------------------------------------------------------------------------------------


def run_weekly_report(
    settings: Settings,
    *,
    now: datetime.datetime | None = None,
    week_of: datetime.date | None = None,
    send_email: bool | None = None,
    session_factory=None,
    sender=None,
) -> WeeklyReport:
    """Build the report, save it, and email it (unless told not to). Raises on a database problem; an email problem is
    raised too, after the file is saved, so a scheduled run shows the failure instead of hiding it."""
    from naukri_agent.database.base import init_db, session_scope

    tz = ZoneInfo(settings.timezone)
    now = now or datetime.datetime.now(datetime.UTC)
    local_now = now.astimezone(tz)
    monday, sunday = week_bounds(local_now.date(), week_of=week_of)
    factory = session_factory or init_db(settings)
    with session_scope(factory) as session:
        report = WeeklyReport(week_start=monday, week_end=sunday, generated_local=local_now)
        report.rows = collect_week(session, settings, monday, sunday)
    write_workbook(report, Path(settings.weekly_report_dir) / report_filename(monday, sunday))
    if settings.weekly_report_email if send_email is None else send_email:
        from naukri_agent.notifications.email import EmailMessage, build_email_sender

        subject, body = email_text(report)
        (sender or build_email_sender(settings)).send(
            EmailMessage(subject=subject, text_body=body, attachments=[str(report.path.resolve())])
        )
    return report
