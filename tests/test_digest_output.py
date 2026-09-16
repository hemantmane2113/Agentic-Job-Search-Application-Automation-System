"""Stage A: digest rendering, file/console email senders, Excel export.

DB is authoritative; Excel is regenerated from it; no secrets are ever
rendered or written.
"""

from __future__ import annotations

import datetime

from openpyxl import load_workbook

from naukri_agent.database.base import session_scope
from naukri_agent.database.models import ApplicationStatus, DailyRunStatus
from naukri_agent.database.repositories import upsert_application_history
from naukri_agent.notifications.email import (
    ConsoleEmailSender,
    FileEmailSender,
    SmtpEmailSender,
    build_email_sender,
)
from naukri_agent.notifications.exceptions import EmailConfigError
from naukri_agent.notifications.render import render_digest
from naukri_agent.recommendations.builder import build_digest
from naukri_agent.reporting.excel import export_workbook

from .digest_fakes import (
    add_job,
    in_memory_factory,
    make_candidate,
    make_run,
    score,
    select_static_resume,
    settings,
)

UTC = datetime.UTC


def _seed_digest(s, cfg):
    cand = make_candidate(s, email="me@example.com")
    run = make_run(s)
    j1 = add_job(s, slug="ds", ext="040926000001", title="Data Scientist", company="Acme",
                 location="Pune", salary_text="12-18 LPA", experience_text="3-6 yrs")
    j2 = add_job(s, slug="mle", ext="040926000002", title="ML Engineer", company="Beta")
    score(s, cand.id, j1.id, overall=88.0)
    score(s, cand.id, j2.id, overall=71.0)
    select_static_resume(s, j1.id, cand.id, resume_id="data_scientist")
    select_static_resume(s, j2.id, cand.id, resume_id="ai_ml_engineer")
    d = build_digest(
        s, candidate_id=cand.id, candidate_email=cand.email,
        scored_job_ids=[j1.id, j2.id], settings=cfg, run_id=run.id,
        now=datetime.datetime(2026, 9, 9, tzinfo=UTC),
    )
    return cand, run, d


def test_render_digest_contains_required_fields_and_no_secrets(tmp_path):
    cfg = settings(tmp_path, smtp_password="SUPER_SECRET_PW", groq_api_key="GROQ_SECRET")
    with session_scope(in_memory_factory()) as s:
        _cand, _run, d = _seed_digest(s, cfg)
        msg = render_digest(d, cfg)

    body = msg.text_body
    assert "09 Sep 2026" in body
    assert "Recommendations: 2" in body
    for token in ("Data Scientist", "Acme", "Pune", "3-6 yrs", "12-18 LPA",
                  "Match score: 88.0", "Recommended resume: data_scientist",
                  "Application status: New", "Freshness: Newly discovered",
                  "Why it matches:", "Important gaps:",
                  "https://www.naukri.com/"):
        assert token in body, token
    assert "SUPER_SECRET_PW" not in body
    assert "GROQ_SECRET" not in body
    assert msg.subject.startswith("[naukri-agent]")


def test_render_digest_shows_previously_applied_only_from_db(tmp_path):
    cfg = settings(tmp_path, recommendation_exclude_if_status=[])  # allow applied through so we can render it
    with session_scope(in_memory_factory()) as s:
        cand = make_candidate(s)
        run = make_run(s)
        j = add_job(s, slug="app", ext="040926000005")
        score(s, cand.id, j.id, overall=90.0)
        upsert_application_history(
            s, j.id, status=ApplicationStatus.APPLIED,
            applied_at=datetime.datetime(2026, 8, 28, tzinfo=UTC),
        )
        d = build_digest(
            s, candidate_id=cand.id, candidate_email=None,
            scored_job_ids=[j.id], settings=cfg, run_id=run.id,
            now=datetime.datetime(2026, 9, 9, tzinfo=UTC),
        )
        body = render_digest(d, cfg).text_body
    assert "Application status: Previously applied — 28 Aug 2026" in body


def test_file_email_sender_writes_and_never_touches_smtp(tmp_path):
    out = tmp_path / "emails"
    sender = FileEmailSender(out)
    from naukri_agent.notifications.email import EmailMessage

    res = sender.send(EmailMessage(subject="s", text_body="hello digest"))
    assert res.sender == "file" and res.status == "written"
    written = list(out.glob("digest_*.txt"))
    assert len(written) == 1 and "hello digest" in written[0].read_text(encoding="utf-8")


def test_build_email_sender_selection(tmp_path):
    assert isinstance(build_email_sender(settings(tmp_path, email_sender="file")), FileEmailSender)
    assert isinstance(build_email_sender(settings(tmp_path, email_sender="console")), ConsoleEmailSender)


def test_smtp_selected_only_when_fully_configured(tmp_path):
    fully_configured = settings(
        tmp_path, email_sender="smtp", smtp_host="smtp.gmail.com", smtp_port=587,
        smtp_username="me@gmail.com", smtp_password="an-app-password", notify_email_to="me@gmail.com",
    )
    sender = build_email_sender(fully_configured)
    assert isinstance(sender, SmtpEmailSender)
    assert sender.host == "smtp.gmail.com" and sender.to_addr == "me@gmail.com"


def test_smtp_missing_credentials_fails_clearly_not_silently():
    for missing_field in ("smtp_host", "smtp_username", "smtp_password", "notify_email_to"):
        overrides = {
            "email_sender": "smtp", "smtp_host": "smtp.gmail.com", "smtp_username": "me@gmail.com",
            "smtp_password": "pw", "notify_email_to": "me@gmail.com", missing_field: "",
        }
        try:
            import tempfile

            with tempfile.TemporaryDirectory() as td:
                from pathlib import Path

                build_email_sender(settings(Path(td), **overrides))
            assert False, f"expected EmailConfigError when {missing_field} is missing"
        except EmailConfigError as exc:
            assert missing_field.upper() in str(exc)


def test_smtp_sender_never_logs_password(monkeypatch, caplog):
    import logging

    sent = {}

    class _FakeSmtp:
        def __init__(self, host, port, timeout=30):
            sent["host"] = host
            sent["port"] = port

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def starttls(self):
            pass

        def login(self, username, password):
            sent["username"] = username
            sent["password"] = password

        def send_message(self, mime_msg):
            sent["subject"] = mime_msg["Subject"]

    monkeypatch.setattr("naukri_agent.notifications.email.smtplib.SMTP", _FakeSmtp)

    from naukri_agent.notifications.email import EmailMessage

    sender = SmtpEmailSender(
        host="smtp.gmail.com", port=587, username="me@gmail.com",
        password="SUPER_SECRET_APP_PASSWORD", to_addr="me@gmail.com",
    )
    with caplog.at_level(logging.DEBUG):
        result = sender.send(EmailMessage(subject="Digest", text_body="body"))

    assert result.sender == "smtp" and result.status == "sent"
    assert sent["password"] == "SUPER_SECRET_APP_PASSWORD"  # really sent to the (fake) server
    assert "SUPER_SECRET_APP_PASSWORD" not in caplog.text  # but never logged


def test_smtp_sender_raises_email_send_error_on_failure_without_leaking_details(monkeypatch):
    import smtplib as smtplib_module

    class _FailingSmtp:
        def __init__(self, host, port, timeout=30):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def starttls(self):
            pass

        def login(self, username, password):
            raise smtplib_module.SMTPAuthenticationError(535, b"5.7.8 Username and Password not accepted, SECRET_TOKEN")

    monkeypatch.setattr("naukri_agent.notifications.email.smtplib.SMTP", _FailingSmtp)

    from naukri_agent.notifications.email import EmailMessage
    from naukri_agent.notifications.exceptions import EmailSendError

    sender = SmtpEmailSender(host="smtp.gmail.com", port=587, username="me@gmail.com",
                              password="pw", to_addr="me@gmail.com")
    try:
        sender.send(EmailMessage(subject="s", text_body="b"))
        assert False, "expected EmailSendError"
    except EmailSendError as exc:
        assert "SECRET_TOKEN" not in str(exc)
        assert "SMTPAuthenticationError" in str(exc)


def test_excel_is_regenerated_entirely_from_db(tmp_path):
    cfg = settings(tmp_path)
    path = tmp_path / "wb.xlsx"
    with session_scope(in_memory_factory()) as s:
        cand = make_candidate(s)
        run = make_run(s)
        run.status = DailyRunStatus.COMPLETED
        run.jobs_discovered = 2
        run.jobs_evaluated = 2
        applied = add_job(s, slug="ap", ext="040926000010", title="Applied Role", company="A")
        seen = add_job(s, slug="se", ext="040926000011", title="Seen Role", company="B")
        score(s, cand.id, applied.id, overall=90.0)
        score(s, cand.id, seen.id, overall=77.0)
        select_static_resume(s, applied.id, cand.id, resume_id="data_scientist")
        upsert_application_history(
            s, applied.id, status=ApplicationStatus.APPLIED, resume_id="data_scientist",
            applied_at=datetime.datetime(2026, 8, 20, tzinfo=UTC), note="applied via portal",
        )
        res1 = export_workbook(s, path, cfg)
        blob1 = path.read_bytes()
        res2 = export_workbook(s, path, cfg)
        blob2 = path.read_bytes()

    assert res1.jobs_rows == 2 and res1.applications_rows == 1 and res1.daily_runs_rows == 1
    assert blob1 == blob2  # idempotent for unchanged DB

    wb = load_workbook(path)
    assert wb.sheetnames == ["Jobs", "Applications", "Daily Runs"]
    jobs = list(wb["Jobs"].iter_rows(values_only=True))
    headers = jobs[0]
    rows = {r[headers.index("Job Title")]: r for r in jobs[1:]}
    ar = rows["Applied Role"]
    assert ar[headers.index("Naukri Job ID")] == "040926000010"
    assert ar[headers.index("Match Score")] == 90.0
    assert ar[headers.index("Application Status")] == "APPLIED"
    assert ar[headers.index("Applied Date")] == "2026-08-20"
    assert ar[headers.index("Resume Used")] == "data_scientist"
    assert ar[headers.index("Notes")] == "applied via portal"
    sr = rows["Seen Role"]
    assert sr[headers.index("Application Status")] == "Not applied"

    apps = list(wb["Applications"].iter_rows(values_only=True))
    assert apps[1][apps[0].index("Job Title")] == "Applied Role"
    runs = list(wb["Daily Runs"].iter_rows(values_only=True))
    rr = runs[1]
    assert rr[runs[0].index("Jobs Discovered")] == 2
    assert rr[runs[0].index("Run Status")] == "COMPLETED"

    # no secret ever lands in a cell
    dumped = str([list(sh.iter_rows(values_only=True)) for sh in wb.worksheets])
    for secret in ("smtp_password", "groq_api_key", "naukri_password", cfg.database_url):
        assert secret not in dumped


def test_excel_regeneration_does_not_preserve_hand_edits(tmp_path):
    cfg = settings(tmp_path)
    path = tmp_path / "wb.xlsx"
    with session_scope(in_memory_factory()) as s:
        cand = make_candidate(s)
        j = add_job(s, slug="x", ext="040926000020")
        score(s, cand.id, j.id, overall=80.0)
        export_workbook(s, path, cfg)

        wb = load_workbook(path)
        wb["Jobs"].cell(row=2, column=20).value = "HAND EDITED NOTE"
        wb.save(path)

        export_workbook(s, path, cfg)  # regenerate

    wb2 = load_workbook(path)
    assert wb2["Jobs"].cell(row=2, column=20).value != "HAND EDITED NOTE"
