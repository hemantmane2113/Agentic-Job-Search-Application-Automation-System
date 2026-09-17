"""Stage A CLI surface: mark-applied / mark-status / applications /
export-excel / doctor, exercised through Click's CliRunner against a
throwaway file database. No browser, no network, no email.
"""

from __future__ import annotations

import datetime
import json
import shutil

import pytest
from click.testing import CliRunner
from openpyxl import load_workbook

import naukri_agent.cli.main as cli_main
from naukri_agent.config import Settings
from naukri_agent.database.base import init_db, session_scope
from naukri_agent.database.models import ApplicationHistory, ApplicationStatus
from naukri_agent.database.repositories import upsert_job, upsert_job_match, upsert_resume_selection
from naukri_agent.jobs.models import JobCreate
from naukri_agent.matching.models import CategoryScore, MatchDecision, MatchResult
from naukri_agent.recommendations.builder import build_digest
from naukri_agent.resume.registry import ResumeMatchVia, ResumeSelectionDecision
from naukri_agent.resume.selector import ResumeSelectionOutcome

EXT = "040926000777"
URL = f"https://www.naukri.com/job-listings-gen-ai-ds-{EXT}"


@pytest.fixture
def env(tmp_path, monkeypatch):
    import naukri_agent

    root = __import__("pathlib").Path(naukri_agent.__file__).resolve().parents[2]
    cfg = root / "config"
    shutil.copy(cfg / "candidate_profile.example.yaml", tmp_path / "cand.yaml")
    shutil.copy(cfg / "master_resume.example.yaml", tmp_path / "resume.yaml")
    shutil.copy(cfg / "resumes.example.yaml", tmp_path / "resumes.yaml")

    settings = Settings(
        _env_file=None,
        database_url=f"sqlite:///{tmp_path / 'cli.db'}",
        candidate_profile_path=str(tmp_path / "cand.yaml"),
        master_resume_path=str(tmp_path / "resume.yaml"),
        resume_registry_path=str(tmp_path / "resumes.yaml"),
        email_output_dir=str(tmp_path / "emails"),
        excel_path=str(tmp_path / "history.xlsx"),
        dry_run=True,
        threshold_review=0,
        threshold_accept=100,
        recommendation_cooldown_days=0,
    )
    monkeypatch.setattr(cli_main, "get_settings", lambda: settings)

    factory = init_db(settings)
    with session_scope(factory) as s:
        from naukri_agent.database.models import Candidate

        cand = Candidate(full_name="T", email="t@example.com", years_experience=4, profile_json="{}")
        s.add(cand)
        s.flush()
        job, _ = upsert_job(
            s,
            JobCreate(
                title="Gen AI DS", company="MNC", location="Pune",
                description="JD", url=URL, salary_text="10-20 LPA", experience_text="2-6 yrs",
            ),
        )
        mr = MatchResult(
            overall_score=82.0, decision=MatchDecision.ACCEPT,
            category_scores={"skills": CategoryScore(points=82, max_points=100)},
            positive_factors=["python"], negative_factors=["salary"],
        )
        upsert_job_match(s, cand.id, job.id, mr)
        upsert_resume_selection(
            s, job.id,
            ResumeSelectionOutcome(
                decision=ResumeSelectionDecision.SELECTED, resume_id="data_scientist",
                file="resumes/data_scientist.pdf", file_hash="abc",
                matched_via=ResumeMatchVia.DETERMINISTIC, reason="role",
            ),
            candidate_id=cand.id,
        )
        cand_id, job_id = cand.id, job.id

    return settings, factory, cand_id, job_id


def _run(*args):
    return CliRunner().invoke(cli_main.cli, list(args))


def test_mark_applied_records_history_and_excludes_from_future_digests(env):
    settings, factory, cand_id, job_id = env

    res = _run("mark-applied", EXT, "--note", "applied on portal")
    assert res.exit_code == 0, res.output
    assert "Recorded APPLIED" in res.output and "exclude" in res.output

    with session_scope(factory) as s:
        row = s.query(ApplicationHistory).one()
        assert row.job_id == job_id
        assert row.status == ApplicationStatus.APPLIED
        assert row.resume_id == "data_scientist"  # pulled from the latest selection
        assert row.notes == "applied on portal"

        # verification point 5: no longer eligible for recommendation
        digest = build_digest(
            s, candidate_id=cand_id, candidate_email="t@example.com",
            scored_job_ids=[job_id], settings=settings, run_id=None,
            now=datetime.datetime(2026, 9, 10, tzinfo=datetime.UTC),
        )
    assert digest.count == 0


def test_mark_applied_accepts_url_and_internal_id_too(env):
    _settings, factory, _cand, job_id = env
    assert _run("mark-applied", URL).exit_code == 0
    # second call via internal id updates the same row, no duplicate
    r2 = _run("mark-applied", str(job_id), "--status", "INTERVIEW")
    assert r2.exit_code == 0 and "Updated INTERVIEW" in r2.output
    with session_scope(factory) as s:
        assert s.query(ApplicationHistory).count() == 1


def test_mark_status_transitions_and_applications_lists_it(env):
    _settings, _factory, _cand, _job = env
    _run("mark-applied", EXT)
    r = _run("mark-status", EXT, "offer")
    assert r.exit_code == 0 and "-> OFFER" in r.output

    listed = _run("applications")
    assert listed.exit_code == 0
    rows = json.loads(listed.output)
    assert len(rows) == 1
    assert rows[0]["status"] == "OFFER"
    assert rows[0]["external_job_id"] == EXT
    assert rows[0]["url"] == URL

    filtered = _run("applications", "--status", "APPLIED")
    assert json.loads(filtered.output) == []


def test_export_excel_regenerates_workbook_from_db(env):
    settings, _factory, _cand, _job = env
    _run("mark-applied", EXT)
    out = settings.excel_path.parent / "cli_export.xlsx"
    res = _run("export-excel", "--path", str(out))
    assert res.exit_code == 0, res.output
    assert out.exists()
    wb = load_workbook(out)
    assert wb.sheetnames == ["Jobs", "Applications", "Daily Runs"]
    jobs = list(wb["Jobs"].iter_rows(values_only=True))
    headers = jobs[0]
    row = jobs[1]
    assert row[headers.index("Naukri Job ID")] == EXT
    assert row[headers.index("Application Status")] == "APPLIED"


def test_doctor_reports_stage_a_checks_ok(env):
    res = _run("doctor")
    # overall exit code may be non-zero in a throwaway env (no LLM creds,
    # resume PDFs absent) — assert the Stage A digest checks themselves pass.
    assert "[OK] Daily recommendation limit" in res.output
    assert "[OK] Recommendation cooldown" in res.output
    assert "[OK] recommendation_exclude_if_status" in res.output
    assert "[OK] Digest tables reachable" in res.output
    assert "application_history=0" in res.output


def test_console_script_entry_point_goes_through_setup_logging_first():
    """Windows Task Scheduler (and any other non-interactive launcher)
    must invoke `naukri-agent ...`, which pyproject.toml wires to
    cli.main:main -- NOT the bare `cli` Click group -- specifically
    because main() calls setup_logging() before cli() runs, so a
    scheduled run's log lines land in LOG_DIR (file, durable) and not
    only on a console nobody is attached to. A regression here (e.g.
    pointing the script at `cli` directly) would silently lose all
    scheduled-run logging."""
    import tomllib
    from pathlib import Path

    import naukri_agent

    root = Path(naukri_agent.__file__).resolve().parents[2]
    pyproject = tomllib.loads((root / "pyproject.toml").read_text(encoding="utf-8"))
    assert pyproject["project"]["scripts"]["naukri-agent"] == "naukri_agent.cli.main:main"

    import inspect

    src = inspect.getsource(cli_main.main)
    assert "setup_logging" in src


def test_run_daily_command_path_never_calls_interactive_input():
    """Structural guard: the production run-daily/discover/recommend
    commands, and everything they import (orchestration + browser
    login, NOT the frozen inspect/inspect-apply tools), must never
    call input() or a Click interactive prompt -- Task Scheduler runs
    with no stdin attached, so any interactive pause would hang the
    scheduled task forever instead of failing cleanly."""
    import inspect

    from naukri_agent.browser import login as login_mod
    from naukri_agent.orchestration import discovery, pipeline

    for mod in (pipeline, discovery, login_mod):
        src = inspect.getsource(mod)
        assert "input(" not in src
        assert "click.prompt" not in src
        assert "click.confirm" not in src
