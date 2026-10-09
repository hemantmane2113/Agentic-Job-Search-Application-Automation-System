"""'Did you apply?' on Telegram for company-website jobs: Applied / Not applying / Later, the 4th Later in a row ignores a
job, the labels in the history, the evening reminder, and nothing is recorded unless you tap."""

from __future__ import annotations

import datetime
import json

import pytest
from openpyxl import load_workbook
from sqlalchemy.exc import OperationalError

from naukri_agent.config import Settings
from naukri_agent.database.base import session_scope
from naukri_agent.database.models import ApplicationEvent, ApplicationHistory, ApplicationStatus, FollowupPrompt
from naukri_agent.database.repositories import (
    LABEL_APPLIED_COMPANY,
    LABEL_APPLIED_DIRECTLY,
    LABEL_IGNORED,
    LABEL_NOT_APPLIED,
    application_label,
    application_status_for_job,
    backfill_apply_labels,
    derive_apply_label,
    excluded_job_ids_by_status,
    followup_candidates,
    record_followup_answer,
    record_job_recommendation,
    upsert_application_history,
)
from naukri_agent.matching.models import MatchDecision
from naukri_agent.orchestration import followup as fu
from naukri_agent.orchestration.followup import FollowupSession, job_card, mark_reminder_handled, reminder_due
from naukri_agent.orchestration.telegram_listener import Listener
from naukri_agent.recommendations.builder import build_digest
from naukri_agent.reporting.excel import export_workbook

from .digest_fakes import add_job, in_memory_factory, make_candidate, make_run, score, select_static_resume, settings

NOW = datetime.datetime(2026, 10, 9, 15, 0, tzinfo=datetime.UTC)  # 20:30 in India


def cfg(tmp_path, **over):
    base = dict(followup_state_file=tmp_path / "followup.json", timezone="Asia/Kolkata", telegram_bot_token="T",
                telegram_chat_id="1", telegram_apply_lock_file=tmp_path / "lock")
    base.update(over)
    return settings(tmp_path, **base)


def seed(s, n=3, *, days_ago=0, apply_type="company_site"):
    """n company-website jobs recommended `days_ago` days before NOW. Returns the ids, newest recommendation first."""
    cand, run = make_candidate(s), make_run(s)
    ids = []
    for i in range(n):
        job = add_job(s, slug=f"j{i}", ext=f"06102650{i:04d}", title=f"role{i}", company=f"Co{i}")
        job.apply_type, job.apply_redirect_url = apply_type, f"https://co{i}.example/jobs/{i}"
        score(s, cand.id, job.id, overall=90.0 - i)
        select_static_resume(s, job.id, cand.id, resume_id="ai_ml_engineer")
        record_job_recommendation(
            s, candidate_id=cand.id, job_id=job.id, daily_run_id=run.id, rank=i + 1, score_at_email=90.0 - i,
            decision_at_email=MatchDecision.ACCEPT, application_status_at_email=ApplicationStatus.NOT_APPLIED,
            recommended_at=NOW - datetime.timedelta(days=days_ago, minutes=i),
        )
        ids.append(job.id)
    return ids


# --- the answers and the labels ------------------------------------------------------------------------------------------


def test_applied_is_recorded_as_applied_through_company_website_with_its_resume():
    with session_scope(in_memory_factory()) as s:
        a, *_ = seed(s)
        assert record_followup_answer(s, a, "applied", ignore_after_later=4, now=NOW) == ("applied", 0)
        row = s.query(ApplicationHistory).one()
        assert row.status == ApplicationStatus.APPLIED and row.apply_label == LABEL_APPLIED_COMPANY
        assert row.source == "telegram_followup" and row.resume_id == "ai_ml_engineer" and row.applied_at is not None
        assert s.query(ApplicationEvent).count() == 1


def test_not_applying_is_recorded_as_not_applied():
    with session_scope(in_memory_factory()) as s:
        a, *_ = seed(s)
        assert record_followup_answer(s, a, "not_applying", ignore_after_later=4)[0] == "not_applying"
        row = s.query(ApplicationHistory).one()
        assert row.status == ApplicationStatus.NOT_APPLYING and row.apply_label == LABEL_NOT_APPLIED and row.applied_at is None


def test_the_fourth_later_in_a_row_turns_the_job_into_ignored_and_not_before():
    with session_scope(in_memory_factory()) as s:
        a, *_ = seed(s)
        got = [record_followup_answer(s, a, "later", ignore_after_later=4) for _ in range(3)]
        assert got == [("later", 1), ("later", 2), ("later", 3)] and s.query(ApplicationHistory).count() == 0
        assert record_followup_answer(s, a, "later", ignore_after_later=4) == ("ignored", 4)
        row = s.query(ApplicationHistory).one()
        assert row.status == ApplicationStatus.IGNORED and row.apply_label == LABEL_IGNORED


def test_an_answer_after_some_laters_settles_the_job():
    with session_scope(in_memory_factory()) as s:
        a, *_ = seed(s)
        record_followup_answer(s, a, "later", ignore_after_later=4)
        record_followup_answer(s, a, "later", ignore_after_later=4)
        assert record_followup_answer(s, a, "applied", ignore_after_later=4)[0] == "applied"
        assert s.query(ApplicationHistory).one().status == ApplicationStatus.APPLIED


def test_an_unknown_answer_is_refused():
    with session_scope(in_memory_factory()) as s:
        a, *_ = seed(s)
        with pytest.raises(ValueError):
            record_followup_answer(s, a, "maybe", ignore_after_later=4)


def test_the_labels_tell_apart_what_the_app_did_and_what_you_did():
    assert derive_apply_label(ApplicationStatus.APPLIED, "agent_auto_apply_telegram") == LABEL_APPLIED_DIRECTLY
    assert derive_apply_label(ApplicationStatus.APPLIED, "manual_cli") == LABEL_APPLIED_COMPANY
    assert derive_apply_label(ApplicationStatus.APPLIED, "telegram_followup") == LABEL_APPLIED_COMPANY
    assert derive_apply_label(ApplicationStatus.NOT_APPLYING, "telegram_followup") == LABEL_NOT_APPLIED
    assert derive_apply_label(ApplicationStatus.IGNORED, "telegram_followup") == LABEL_IGNORED
    assert derive_apply_label(ApplicationStatus.INTERVIEW, "manual_cli") is None


def test_marking_by_hand_and_the_apps_own_applications_get_their_labels():
    with session_scope(in_memory_factory()) as s:
        a, b, c = seed(s)
        by_hand, _ = upsert_application_history(s, a)  # `mark-applied`
        by_app, _ = upsert_application_history(s, b, source="agent_auto_apply_telegram")
        assert by_hand.apply_label == LABEL_APPLIED_COMPANY and by_app.apply_label == LABEL_APPLIED_DIRECTLY
        later, _ = upsert_application_history(s, a, status=ApplicationStatus.INTERVIEW)
        assert later.apply_label == LABEL_APPLIED_COMPANY  # an interview does not erase how it was applied


def test_rows_from_before_labels_existed_are_labelled_on_read_and_by_the_backfill():
    with session_scope(in_memory_factory()) as s:
        a, b, _c = seed(s)
        upsert_application_history(s, a, source="agent_auto_apply_telegram")
        upsert_application_history(s, b)
        s.query(ApplicationHistory).update({ApplicationHistory.apply_label: None})
        rows = s.query(ApplicationHistory).order_by(ApplicationHistory.job_id).all()
        assert [application_label(r) for r in rows] == [LABEL_APPLIED_DIRECTLY, LABEL_APPLIED_COMPANY]
        assert backfill_apply_labels(s) == 2 and backfill_apply_labels(s) == 0
        assert all(r.apply_label for r in rows)


def test_not_applying_and_ignored_are_excluded_from_future_digests_by_default():
    c = Settings(_env_file=None)
    assert {"NOT_APPLYING", "IGNORED"} <= set(c.recommendation_exclude_if_status)
    with session_scope(in_memory_factory()) as s:
        a, b, _c = seed(s)
        record_followup_answer(s, a, "not_applying", ignore_after_later=4)
        for _ in range(4):
            record_followup_answer(s, b, "later", ignore_after_later=4)
        assert excluded_job_ids_by_status(s, c.recommendation_exclude_if_status) == {a, b}


def test_a_job_you_said_not_applying_to_never_comes_back_in_the_digest(tmp_path):
    cfg_ = cfg(tmp_path)
    with session_scope(in_memory_factory()) as s:
        cand, run = make_candidate(s), make_run(s)
        job = add_job(s, slug="x", ext="061026509999", title="role", company="Co")
        job.apply_type = "company_site"
        score(s, cand.id, job.id, overall=92.0)
        select_static_resume(s, job.id, cand.id)
        record_followup_answer(s, job.id, "not_applying", ignore_after_later=4)
        d = build_digest(s, candidate_id=cand.id, candidate_email="c@example.com", scored_job_ids=[job.id],
                         settings=cfg_, run_id=run.id, now=NOW)
        assert d.count == 0


# --- which jobs get asked about -----------------------------------------------------------------------------------------


def test_only_recent_company_website_jobs_without_an_answer_are_asked_about():
    with session_scope(in_memory_factory()) as s:
        a, b, c = seed(s)
        record_followup_answer(s, c, "applied", ignore_after_later=4)  # answered
        got = followup_candidates(s, NOW, lookback_days=7)
        assert [j["job_id"] for j in got] == [a, b]
        assert got[0]["apply_redirect_url"] == "https://co0.example/jobs/0" and got[0]["resume_id"] == "ai_ml_engineer"


def test_naukri_apply_jobs_and_old_digests_are_not_asked_about():
    with session_scope(in_memory_factory()) as s:
        seed(s, 2, apply_type="native")
        assert followup_candidates(s, NOW, 7) == []
    with session_scope(in_memory_factory()) as s:
        seed(s, 2, days_ago=9)
        assert followup_candidates(s, NOW, 7) == []


def test_a_later_job_is_asked_again_with_its_count_and_a_cap_on_how_many_are_asked():
    with session_scope(in_memory_factory()) as s:
        a, *_ = seed(s, 5)
        record_followup_answer(s, a, "later", ignore_after_later=4)
        got = followup_candidates(s, NOW, 7, limit=3)
        assert len(got) == 3 and {j["job_id"]: j["later_count"] for j in got}[a] == 1


# --- the questions on Telegram -----------------------------------------------------------------------------------------


class Channel:
    def __init__(self, replies):
        self.replies, self.sent, self.asked = list(replies), [], []

    def send(self, text, **_):
        self.sent.append(text)

    def ask_choice(self, text, options, timeout_s):
        self.asked.append((text, list(options), timeout_s))
        return self.replies.pop(0) if self.replies else None


def run_session(tmp_path, factory, replies, **kw):
    ch = Channel(replies)
    status = FollowupSession(cfg(tmp_path), ch, factory=factory).run(now=NOW, **kw)
    return status, ch


def test_each_job_is_asked_with_the_three_buttons_and_answers_are_recorded(tmp_path):
    factory = in_memory_factory()
    with session_scope(factory) as s:
        a, b, c = seed(s)
    status, ch = run_session(tmp_path, factory, ["Applied", "Not applying", "Later"])
    assert status == "done" and len(ch.asked) == 3
    assert all(opts == ["Applied", "Not applying", "Later"] and t == 15 * 60 for _x, opts, t in ch.asked)
    assert "role0" in ch.asked[0][0] and "Direct apply link: https://co0.example/jobs/0" in ch.asked[0][0]
    with session_scope(factory) as s:
        assert application_status_for_job(s, a) == ApplicationStatus.APPLIED
        assert application_status_for_job(s, b) == ApplicationStatus.NOT_APPLYING
        assert application_status_for_job(s, c) == ApplicationStatus.NOT_APPLIED  # Later records nothing yet
        assert s.query(FollowupPrompt).one().later_count == 1
    said = "\n".join(ch.sent)
    assert "Recorded: applied through company website" in said and "Recorded: not applied" in said
    assert "Applied: 1, not applying: 1, later: 1, ignored: 0" in said


def test_no_answer_stops_and_changes_nothing_for_the_rest(tmp_path):
    factory = in_memory_factory()
    with session_scope(factory) as s:
        seed(s)
    status, ch = run_session(tmp_path, factory, ["Applied"])  # the 2nd question gets no tap
    assert len(ch.asked) == 2 and "No answer, so I stopped here" in "\n".join(ch.sent)
    with session_scope(factory) as s:
        assert s.query(ApplicationHistory).count() == 1


def test_stop_ends_the_questions(tmp_path):
    factory = in_memory_factory()
    with session_scope(factory) as s:
        seed(s)
    _status, ch = run_session(tmp_path, factory, ["/stop"])
    assert len(ch.asked) == 1 and "Stopped" in "\n".join(ch.sent)
    with session_scope(factory) as s:
        assert s.query(ApplicationHistory).count() == 0


def test_the_fourth_later_over_four_sessions_ignores_the_job_and_the_third_warns(tmp_path):
    factory = in_memory_factory()
    with session_scope(factory) as s:
        (a,) = seed(s, 1)
    messages = []
    for _ in range(4):
        _status, ch = run_session(tmp_path, factory, ["Later"])
        messages.append("\n".join(ch.sent))
    assert "Okay, I will ask again later." in messages[0] and "Last time I will ask" in messages[2]
    assert "Recorded: ignored" in messages[3]
    with session_scope(factory) as s:
        assert application_status_for_job(s, a) == ApplicationStatus.IGNORED
        assert followup_candidates(s, NOW, 7) == []
    status, ch = run_session(tmp_path, factory, [])  # nothing left to ask
    assert status == "empty" and "Nothing is waiting" in ch.sent[0]


def test_when_nothing_is_waiting_the_reminder_stays_silent_but_the_command_answers(tmp_path):
    factory = in_memory_factory()
    status, ch = run_session(tmp_path, factory, [], quiet_if_empty=True)
    assert status == "empty" and ch.sent == []


def test_a_busy_database_is_reported_and_asks_nothing(tmp_path):
    def busy():
        raise OperationalError("x", {}, Exception("database is locked"))

    status, ch = run_session(tmp_path, busy, [])
    assert status == "busy" and "daily run is probably still going" in ch.sent[0] and ch.asked == []
    status, ch = run_session(tmp_path, busy, [], quiet_if_empty=True)
    assert status == "busy" and ch.sent == []


def test_the_card_shows_what_helps_decide():
    job = {"title": "AI Engineer", "company": "LG", "location": "Bengaluru", "url": "https://naukri.example/j",
           "apply_redirect_url": "https://lg.example/apply", "resume_id": "ai_ml_engineer", "later_count": 2}
    card = job_card(2, 5, job, 4)
    assert "(2 of 5)" in card and "https://lg.example/apply" in card and "Later 2 time(s)" in card and "ai_ml_engineer" in card
    assert "Direct apply link" not in job_card(1, 1, {**job, "apply_redirect_url": None}, 4)


# --- the evening reminder --------------------------------------------------------------------------------------------------


def at(hh, mm, day=9):  # an instant at hh:mm India time on 2026-10-<day>
    return datetime.datetime(2026, 10, day, hh, mm, tzinfo=datetime.timezone(datetime.timedelta(hours=5, minutes=30))).timestamp()


def test_the_reminder_is_due_from_eight_pm_once_a_day(tmp_path):
    c = cfg(tmp_path)
    assert not reminder_due(c, at(19, 59)) and reminder_due(c, at(20, 0)) and reminder_due(c, at(23, 30))
    mark_reminder_handled(c, at(20, 1))
    assert not reminder_due(c, at(20, 5)) and not reminder_due(c, at(23, 59))
    assert reminder_due(c, at(20, 0, day=10))  # tomorrow evening it is due again
    assert json.loads(c.followup_state_file.read_text()) == {"last_reminder_date": "2026-10-09"}


def test_the_reminder_can_be_turned_off_and_its_time_changed(tmp_path):
    assert not reminder_due(cfg(tmp_path, followup_reminder_enabled=False), at(21, 0))
    c = cfg(tmp_path, followup_reminder_time="18:30")
    assert reminder_due(c, at(18, 30)) and not reminder_due(c, at(18, 29))
    with pytest.raises(ValueError):
        Settings(_env_file=None, followup_reminder_time="8pm")


class FakeChannel:
    def __init__(self):
        self.sent = []

    def poll_commands(self, seconds):
        return []

    def send(self, text, **_):
        self.sent.append(text)


def listener(tmp_path, results, clock_value, **over):
    calls = []

    def fake_followup(settings, channel, quiet_if_empty=False):
        calls.append(quiet_if_empty)
        return results.pop(0) if results else "done"

    now = {"t": clock_value}
    lis = Listener(cfg(tmp_path, **over), FakeChannel(), spawn=lambda s: (0, ""), summary=lambda s, n: (0, 4, 4),
                   clock=lambda: now["t"], sleep=lambda s: None, lock_held=lambda p: False, followup=fake_followup)
    return lis, calls, now


def test_at_eight_pm_the_listener_asks_once_and_then_leaves_the_day_alone(tmp_path):
    lis, calls, _now = listener(tmp_path, ["done"], at(20, 5))
    lis.step()
    lis.step()
    assert calls == [True]  # asked once, quietly, and not again that day


def test_before_eight_pm_it_does_not_ask(tmp_path):
    lis, calls, _now = listener(tmp_path, [], at(19, 0))
    lis.step()
    assert calls == []


def test_a_busy_database_at_eight_pm_is_retried_ten_minutes_later_not_every_poll(tmp_path):
    lis, calls, now = listener(tmp_path, ["busy", "done"], at(20, 0))
    lis.step()
    lis.step()  # still inside the ten minutes
    assert calls == [True]
    now["t"] = at(20, 11)
    lis.step()
    assert calls == [True, True]
    lis.step()
    assert calls == [True, True]  # handled for the day


def test_an_empty_evening_is_handled_so_it_is_not_checked_again_all_night(tmp_path):
    lis, calls, _now = listener(tmp_path, ["empty"], at(20, 0))
    lis.step()
    lis.step()
    assert calls == [True]


def test_the_applied_command_asks_on_request(tmp_path):
    lis, calls, _now = listener(tmp_path, [], at(12, 0))
    lis.handle("/applied")
    assert calls == [False]
    assert "/applied" in lis.ch.sent[-1] if lis.ch.sent else True
    lis.handle("/help")
    assert "/applied" in lis.ch.sent[-1]


def test_the_follow_up_module_has_no_way_to_apply_or_open_naukri():
    import inspect

    src = inspect.getsource(fu)
    for word in ("auto_apply_runner", "apply_workflow", "NaukriClient", "BrowserManager", "prepare_application"):
        assert word not in src


# --- where the label shows -----------------------------------------------------------------------------------------------------


def test_the_excel_applications_sheet_has_a_label_column(tmp_path):
    with session_scope(in_memory_factory()) as s:
        a, b, _c = seed(s)
        upsert_application_history(s, a, source="agent_auto_apply_telegram")
        record_followup_answer(s, b, "not_applying", ignore_after_later=4)
        out = tmp_path / "h.xlsx"
        export_workbook(s, out)
    ws = load_workbook(out)["Applications"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    assert "Label" in header
    assert {r[header.index("Label")] for r in rows[1:]} == {LABEL_APPLIED_DIRECTLY, LABEL_NOT_APPLIED}


# --- the real Telegram channel end to end (simulated Telegram) -------------------------------------------------------------


def test_applied_command_through_the_real_channel_asks_with_buttons_and_records_the_tap(tmp_path):
    from .test_telegram import World, button, text

    factory = in_memory_factory()
    with session_scope(factory) as s:
        a, *_ = seed(s, 2)
    # batch 1: the "/applied" message; then (while asking) nothing waiting, a tap on "Applied", and silence after that
    w = World([[text("/applied")], [], [button("opt:0")]])
    c = cfg(tmp_path, telegram_chat_id="4242")
    lis = Listener(c, w.channel(), spawn=lambda s_: (0, ""), summary=lambda s_, n: (0, 4, 4),
                   clock=lambda: NOW.timestamp(), sleep=lambda s_: None, lock_held=lambda p: False,
                   followup=lambda settings, channel, quiet_if_empty=False: FollowupSession(settings, channel, factory=factory).run(
                       quiet_if_empty=quiet_if_empty, now=NOW))
    lis.step()
    sent = [p["text"] for p in w.sent()]
    assert sent[0].startswith("2 company-website job(s) are waiting")
    cards = [p for p in w.sent() if "reply_markup" in p]
    assert cards and [b["text"] for b in cards[0]["reply_markup"]["inline_keyboard"][0]] == ["Applied"]
    assert [row[0]["text"] for row in cards[0]["reply_markup"]["inline_keyboard"]] == ["Applied", "Not applying", "Later"]
    with session_scope(factory) as s:
        assert application_status_for_job(s, a) == ApplicationStatus.APPLIED


# --- a dropped connection while waiting for your tap (live 2026-10-09) --------------------------------------------------------


def flaky_world(fail_polls):
    """Telegram that drops the connection on the Nth getUpdates calls listed in fail_polls (1-based)."""
    from naukri_agent.notifications.telegram import TelegramChannel, TelegramError

    from .test_telegram import World, button

    w = World([[], [], [button("opt:2")]])
    state = {"n": 0}
    real = w.transport

    def transport(method, params, timeout):
        if method == "getUpdates":
            state["n"] += 1
            if state["n"] in fail_polls:
                raise TelegramError("getUpdates failed: TimeoutError")
        return real(method, params, timeout)

    sleeps = []
    return w, TelegramChannel("T", "4242", transport=transport, clock=w.clock, sleep=sleeps.append), sleeps


def test_a_dropped_connection_while_waiting_for_a_tap_does_not_lose_the_question():
    w, ch, sleeps = flaky_world({2, 3})  # drain finds nothing; the next two polls drop; then the tap arrives
    assert ch.ask_choice("Did you apply?", ["Applied", "Not applying", "Later"], 600) == "Later"
    assert len(sleeps) == 2  # it paused and listened again each time


def test_a_connection_that_stays_down_eventually_gives_up_instead_of_looping_forever():
    from naukri_agent.notifications.telegram import TelegramChannel, TelegramError

    from .test_telegram import World

    w = World([[]])
    calls = {"n": 0}

    def transport(method, params, timeout):
        if method == "getUpdates" and calls["n"] > 0:
            raise TelegramError("getUpdates failed: URLError/gaierror")
        calls["n"] += 1
        return w.transport(method, params, timeout)

    ch = TelegramChannel("T", "4242", transport=transport, clock=w.clock, sleep=lambda s: None)
    with pytest.raises(TelegramError):
        ch.ask_choice("Q", ["a", "b"], 600)


def test_a_tap_on_an_expired_question_is_answered_not_left_spinning():
    from .test_telegram import World, button, text

    w = World([[button("opt:0"), text("/applied")]])
    got = w.channel().poll_commands(0)
    assert got == [("/applied", None)]  # the tap is not a command
    alerts = [p for m, p in w.calls if m == "answerCallbackQuery"]
    assert len(alerts) == 1 and alerts[0]["show_alert"] and "expired" in alerts[0]["text"]


def test_a_tap_from_a_stranger_is_ignored_silently():
    from .test_telegram import World, button

    w = World([[button("opt:0", chat="999")]])
    assert w.channel().poll_commands(0) == []
    assert not [1 for m, _p in w.calls if m == "answerCallbackQuery"]


def test_the_question_session_stops_cleanly_when_telegram_stays_down(tmp_path):
    from naukri_agent.notifications.telegram import TelegramError

    class Down(Channel):
        def ask_choice(self, text, options, timeout_s):
            raise TelegramError("getUpdates failed: URLError/gaierror")

    factory = in_memory_factory()
    with session_scope(factory) as s:
        seed(s)
    ch = Down([])
    assert FollowupSession(cfg(tmp_path), ch, factory=factory).run(now=NOW) == "done"
    with session_scope(factory) as s:
        assert s.query(ApplicationHistory).count() == 0


def test_one_command_that_blows_up_does_not_end_phone_control(tmp_path):
    def boom(settings, channel, quiet_if_empty=False):
        raise RuntimeError("bug")

    class Ch(FakeChannel):
        def __init__(self):
            super().__init__()
            self.polls = 0

        def poll_commands(self, seconds):
            self.polls += 1
            return [("/applied", None)] if self.polls == 1 else []

    ch = Ch()
    sleeps = []
    lis = Listener(cfg(tmp_path), ch, spawn=lambda s: (0, ""), summary=lambda s, n: (0, 4, 4), clock=lambda: at(12, 0),
                   sleep=sleeps.append, lock_held=lambda p: False, followup=boom)
    lis.run(should_stop=lambda: ch.polls >= 3)
    assert ch.polls >= 3 and sleeps  # it carried on after the failure
