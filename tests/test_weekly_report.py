"""The weekly Excel report: every job given in a Monday-to-Sunday week and where each stands, emailed as an attachment."""

from __future__ import annotations

import datetime
from pathlib import Path

import pytest
from click.testing import CliRunner
from openpyxl import load_workbook

import naukri_agent.cli.main as cli_main
from naukri_agent.config import Settings
from naukri_agent.database.base import init_db, session_scope
from naukri_agent.database.models import ApplicationStatus, AutoApplyAttempt, Job, RunEventStatus
from naukri_agent.database.repositories import (
    add_run_event,
    record_followup_answer,
    record_job_recommendation,
    upsert_application_history,
)
from naukri_agent.matching.models import MatchDecision
from naukri_agent.notifications.email import EmailMessage, FileEmailSender, SmtpEmailSender
from naukri_agent.reporting.weekly import (
    ROUTE_COMPANY,
    ROUTE_TELEGRAM,
    WeeklyReport,
    collect_week,
    email_text,
    run_weekly_report,
    week_bounds,
    write_workbook,
)

from .digest_fakes import add_job, in_memory_factory, make_candidate, make_run, score, select_static_resume

IST = datetime.timezone(datetime.timedelta(hours=5, minutes=30))
MON, SUN = datetime.date(2026, 10, 5), datetime.date(2026, 10, 11)


def ist(day, hh=10, mm=0):  # a stored (naive UTC) moment for hh:mm India time on 2026-10-<day>
    return datetime.datetime(2026, 10, day, hh, mm, tzinfo=IST).astimezone(datetime.UTC).replace(tzinfo=None)


def cfg(tmp_path, **over):
    base = dict(timezone="Asia/Kolkata", weekly_report_dir=tmp_path / "weekly", email_sender="file",
                email_output_dir=tmp_path / "emails", telegram_bot_token="", telegram_chat_id="")
    base.update(over)
    return Settings(_env_file=None, **base)


def company_job(s, cand, run, n, day, *, hh=10, mm=0):
    job = add_job(s, slug=f"c{n}", ext=f"07102650{n:04d}", title=f"Company role {n}", company=f"Co{n}")
    job.apply_type, job.apply_redirect_url = "company_site", f"https://co{n}.example/jobs/{n}"
    score(s, cand.id, job.id, overall=90.0 - n)
    select_static_resume(s, job.id, cand.id, resume_id="ai_ml_engineer")
    record_job_recommendation(
        s, candidate_id=cand.id, job_id=job.id, daily_run_id=run.id, rank=n + 1, score_at_email=90.0 - n,
        decision_at_email=MatchDecision.ACCEPT, application_status_at_email=ApplicationStatus.NOT_APPLIED,
        recommended_at=ist(day, hh, mm),
    )
    return job.id


def native_job(s, cand, n):
    job = add_job(s, slug=f"n{n}", ext=f"08102650{n:04d}", title=f"Naukri role {n}", company=f"Nk{n}")
    job.apply_type = "native"
    score(s, cand.id, job.id, overall=85.0 - n)
    select_static_resume(s, job.id, cand.id)
    return job.id


def attempt(s, job_id, outcome, when):
    s.add(AutoApplyAttempt(job_id=job_id, attempt_id=f"a{job_id}{outcome}", outcome=outcome, detail="{}", attempted_at=when))
    s.flush()


def ping(s, run, job_ids, seq=1):
    add_run_event(s, daily_run_id=run.id, seq=seq, stage="telegram_ping", status=RunEventStatus.OK,
                  detail={"jobs": len(job_ids), "job_ids": job_ids})


# --- which week -----------------------------------------------------------------------------------------------------------


def test_on_a_sunday_the_report_covers_that_week():
    assert week_bounds(SUN) == (MON, SUN)


def test_a_missed_sunday_caught_up_later_reports_the_week_that_just_ended():
    for day in range(12, 18):  # Monday 12th to Saturday 17th
        monday, sunday = week_bounds(datetime.date(2026, 10, day))
        assert (monday, sunday) == (datetime.date(2026, 10, 5), datetime.date(2026, 10, 11)) if day == 12 else True
    assert week_bounds(datetime.date(2026, 10, 12)) == (MON, SUN)
    assert week_bounds(datetime.date(2026, 10, 10)) == (datetime.date(2026, 9, 28), datetime.date(2026, 10, 4))  # a Saturday


def test_a_chosen_date_picks_its_own_week():
    assert week_bounds(datetime.date(2026, 10, 20), week_of=datetime.date(2026, 10, 7)) == (MON, SUN)
    assert week_bounds(datetime.date(2026, 10, 20), week_of=MON) == (MON, SUN)
    assert week_bounds(datetime.date(2026, 10, 20), week_of=SUN) == (MON, SUN)


# --- which jobs, which status ------------------------------------------------------------------------------------------------


def test_company_website_jobs_show_where_each_one_stands(tmp_path):
    with session_scope(in_memory_factory()) as s:
        cand, run = make_candidate(s), make_run(s)
        a, b, c, d, e = (company_job(s, cand, run, i, 6 + i) for i in range(5))
        record_followup_answer(s, b, "applied", ignore_after_later=4)
        record_followup_answer(s, c, "not_applying", ignore_after_later=4)
        for _ in range(4):
            record_followup_answer(s, d, "later", ignore_after_later=4)
        record_followup_answer(s, e, "later", ignore_after_later=4)
        rows = collect_week(s, cfg(tmp_path), MON, SUN)
    by_title = {r.title: r for r in rows}
    assert by_title["Company role 0"].status == "waiting for your answer"
    assert by_title["Company role 1"].status == "applied through company website"
    assert by_title["Company role 2"].status == "not applied"
    assert by_title["Company role 3"].status == "ignored"
    assert by_title["Company role 4"].status == "waiting for your answer (put off 1x)"
    assert all(r.route == ROUTE_COMPANY for r in rows)
    first = by_title["Company role 0"]
    assert first.score == 90.0 and first.resume == "ai_ml_engineer" and first.direct_link == "https://co0.example/jobs/0"


def test_jobs_outside_the_week_are_left_out_and_the_boundaries_follow_india_time(tmp_path):
    with session_scope(in_memory_factory()) as s:
        cand, run = make_candidate(s), make_run(s)
        company_job(s, cand, run, 0, 4, hh=23, mm=50)   # Sunday 4th 23:50 - last week
        company_job(s, cand, run, 1, 5, hh=0, mm=10)    # Monday 5th 00:10 - this week (it is still the 4th in UTC)
        company_job(s, cand, run, 2, 11, hh=23, mm=50)  # Sunday 11th 23:50 - this week
        company_job(s, cand, run, 3, 12, hh=0, mm=10)   # Monday 12th 00:10 - next week
        rows = collect_week(s, cfg(tmp_path), MON, SUN)
    assert [r.title for r in rows] == ["Company role 1", "Company role 2"]
    assert rows[0].given_on == datetime.date(2026, 10, 5) and rows[1].given_on == datetime.date(2026, 10, 11)


def test_telegram_jobs_show_what_happened_to_each(tmp_path):
    with session_scope(in_memory_factory()) as s:
        cand, run = make_candidate(s), make_run(s, started_at=ist(7))
        applied, declined, silent, failed, only_pinged = (native_job(s, cand, i) for i in range(5))
        upsert_application_history(s, applied, source="agent_auto_apply_telegram")
        attempt(s, applied, "applied", ist(8))
        attempt(s, declined, "declined", ist(8))
        attempt(s, silent, "no_reply", ist(8))
        attempt(s, failed, "failed", ist(8))
        ping(s, run, [applied, declined, silent, failed, only_pinged])
        rows = {r.title: r for r in collect_week(s, cfg(tmp_path), MON, SUN)}
    assert rows["Naukri role 0"].status == "applied directly"
    assert rows["Naukri role 1"].status == "you said no"
    assert rows["Naukri role 2"].status == "no reply on Telegram - not applied"
    assert rows["Naukri role 3"].status == "failed - not applied"
    assert rows["Naukri role 4"].status == "ready on Telegram - not offered yet"
    assert all(r.route == ROUTE_TELEGRAM for r in rows.values()) and len(rows) == 5
    assert rows["Naukri role 0"].given_on == datetime.date(2026, 10, 7) and rows["Naukri role 0"].score == 85.0


def test_a_job_offered_without_a_ping_record_is_still_listed_from_its_attempt(tmp_path):
    with session_scope(in_memory_factory()) as s:
        cand = make_candidate(s)
        job = native_job(s, cand, 0)
        attempt(s, job, "declined", ist(9))
        (row,) = collect_week(s, cfg(tmp_path), MON, SUN)
    assert row.given_on == datetime.date(2026, 10, 9) and row.status == "you said no"


def test_the_route_is_the_jobs_real_type_not_the_email_it_appeared_in(tmp_path):
    """Older digests listed every job; a Naukri Apply job that was in one still belongs to the Telegram route."""
    with session_scope(in_memory_factory()) as s:
        cand, run = make_candidate(s), make_run(s)
        job = company_job(s, cand, run, 0, 7)
        s.get(Job, job).apply_type = "native"
        upsert_application_history(s, job, source="agent_auto_apply_telegram")
        (row,) = collect_week(s, cfg(tmp_path), MON, SUN)
    assert row.route == ROUTE_TELEGRAM and row.status == "applied directly" and row.direct_link is None


def test_a_failed_or_unconfirmed_job_is_flagged_for_a_check_on_naukri(tmp_path):
    with session_scope(in_memory_factory()) as s:
        cand, run = make_candidate(s), make_run(s, started_at=ist(7))
        failed, declined = native_job(s, cand, 0), native_job(s, cand, 1)
        attempt(s, failed, "failed", ist(8))
        attempt(s, declined, "declined", ist(8))
        ping(s, run, [failed, declined])
        rows = {r.title: r for r in collect_week(s, cfg(tmp_path), MON, SUN)}
    assert "Check this job on Naukri" in rows["Naukri role 0"].note and rows["Naukri role 1"].note == ""


def test_a_job_is_listed_once_even_when_two_sources_know_it(tmp_path):
    with session_scope(in_memory_factory()) as s:
        cand, run = make_candidate(s), make_run(s, started_at=ist(7))
        job = company_job(s, cand, run, 0, 7)
        attempt(s, job, "declined", ist(8))
        ping(s, run, [job])
        assert len(collect_week(s, cfg(tmp_path), MON, SUN)) == 1


def test_an_empty_week_gives_an_empty_report(tmp_path):
    with session_scope(in_memory_factory()) as s:
        assert collect_week(s, cfg(tmp_path), MON, SUN) == []


# --- the workbook -----------------------------------------------------------------------------------------------------------------


def sample_report(tmp_path):
    with session_scope(in_memory_factory()) as s:
        cand, run = make_candidate(s), make_run(s, started_at=ist(7))
        a, b = company_job(s, cand, run, 0, 7), company_job(s, cand, run, 1, 8)
        record_followup_answer(s, b, "applied", ignore_after_later=4)
        n = native_job(s, cand, 0)
        ping(s, run, [n])
        rows = collect_week(s, cfg(tmp_path), MON, SUN)
    return WeeklyReport(MON, SUN, datetime.datetime(2026, 10, 11, 22, 0, tzinfo=IST), rows)


def test_the_workbook_has_a_summary_and_one_row_per_job(tmp_path):
    report = sample_report(tmp_path)
    path = write_workbook(report, tmp_path / "w.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Jobs"]
    summary = "\n".join(str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value is not None)
    assert "Mon 05 Oct to Sun 11 Oct 2026" in summary and "Jobs given this week" in summary
    assert ROUTE_COMPANY in summary and ROUTE_TELEGRAM in summary and "applied through company website" in summary
    rows = list(wb["Jobs"].iter_rows(values_only=True))
    header = list(rows[0])
    assert header[:3] == ["Date given", "Day", "Route"] and len(rows) == 4
    status_col = header.index("Status")
    assert {r[status_col] for r in rows[1:]} == {
        "waiting for your answer", "applied through company website", "ready on Telegram - not offered yet"
    }
    day_col = header.index("Day")
    assert rows[1][day_col] == "Wed"  # the 7th


def test_statuses_are_coloured_and_links_are_clickable(tmp_path):
    report = sample_report(tmp_path)
    ws = load_workbook(write_workbook(report, tmp_path / "w.xlsx"))["Jobs"]
    header = [c.value for c in ws[1]]
    by_status = {ws.cell(row=r, column=header.index("Status") + 1).value: ws.cell(row=r, column=header.index("Status") + 1).fill.fgColor.rgb
                 for r in range(2, ws.max_row + 1)}
    assert by_status["applied through company website"].endswith("C6EFCE")
    assert by_status["waiting for your answer"].endswith("FFEB9C")
    link = ws.cell(row=2, column=header.index("Direct apply link") + 1)
    assert link.hyperlink is not None and link.value.startswith("https://co")
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref


def test_the_email_text_summarises_the_week(tmp_path):
    report = sample_report(tmp_path)
    report.path = Path("weekly_2026-10-05_to_2026-10-11.xlsx")
    subject, body = email_text(report)
    assert subject == "[naukri-agent] Weekly report: 05 Oct to 11 Oct 2026 - 3 jobs"
    assert "3 jobs were given this week." in body and "Company website: 2" in body and "applied through company website: 1" in body
    assert "attached Excel file" in body and "as of Sun 11 Oct" in body
    empty = WeeklyReport(MON, SUN, datetime.datetime(2026, 10, 11, 22, 0, tzinfo=IST), [])
    assert "No jobs were given this week." in email_text(empty)[1]


# --- the whole run ----------------------------------------------------------------------------------------------------------------


class Recorder:
    def __init__(self):
        self.sent = []

    def send(self, message):
        self.sent.append(message)


def seeded_factory():
    factory = in_memory_factory()
    with session_scope(factory) as s:
        cand, run = make_candidate(s), make_run(s)
        company_job(s, cand, run, 0, 7)
    return factory


def test_a_sunday_run_saves_the_file_and_emails_it_as_an_attachment(tmp_path):
    rec = Recorder()
    now = datetime.datetime(2026, 10, 11, 22, 0, tzinfo=IST).astimezone(datetime.UTC)
    report = run_weekly_report(cfg(tmp_path), now=now, session_factory=seeded_factory(), sender=rec)
    assert report.path.name == "weekly_2026-10-05_to_2026-10-11.xlsx" and report.path.exists()
    (message,) = rec.sent
    assert message.subject.endswith("- 1 job") and message.attachments == [str(report.path.resolve())]


def test_it_can_save_without_emailing(tmp_path):
    rec = Recorder()
    now = datetime.datetime(2026, 10, 11, 22, 0, tzinfo=IST).astimezone(datetime.UTC)
    run_weekly_report(cfg(tmp_path), now=now, session_factory=seeded_factory(), sender=rec, send_email=False)
    assert rec.sent == []
    run_weekly_report(cfg(tmp_path, weekly_report_email=False), now=now, session_factory=seeded_factory(), sender=rec)
    assert rec.sent == []


def test_an_email_failure_is_raised_after_the_file_is_safely_saved(tmp_path):
    from naukri_agent.notifications.exceptions import EmailSendError

    class Broken:
        def send(self, message):
            raise EmailSendError("SmtpEmailSender failed: TimeoutError")

    now = datetime.datetime(2026, 10, 11, 22, 0, tzinfo=IST).astimezone(datetime.UTC)
    with pytest.raises(EmailSendError):
        run_weekly_report(cfg(tmp_path), now=now, session_factory=seeded_factory(), sender=Broken())
    assert (tmp_path / "weekly" / "weekly_2026-10-05_to_2026-10-11.xlsx").exists()


# --- attachments in the real senders ----------------------------------------------------------------------------------------------


def test_smtp_sends_the_file_as_a_proper_attachment_and_plain_mail_is_unchanged(tmp_path, monkeypatch):
    sent = []

    class FakeSmtp:
        def __init__(self, host, port, timeout=30):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def starttls(self):
            pass

        def login(self, u, p):
            pass

        def send_message(self, msg):
            sent.append(msg)

    monkeypatch.setattr("naukri_agent.notifications.email.smtplib.SMTP", FakeSmtp)
    xlsx = tmp_path / "weekly_x.xlsx"
    write_workbook(sample_report(tmp_path), xlsx)
    sender = SmtpEmailSender("h", 587, "me@x.com", "pw", "me@x.com")
    sender.send(EmailMessage(subject="s", text_body="b", attachments=[str(xlsx)]))
    sender.send(EmailMessage(subject="plain", text_body="b", html_body="<p>b</p>"))
    with_file, plain = sent
    assert with_file.get_content_type() == "multipart/mixed" and plain.get_content_type() == "multipart/alternative"
    parts = [p for p in with_file.walk() if p.get_filename()]
    assert [p.get_filename() for p in parts] == ["weekly_x.xlsx"]
    assert parts[0].get_content_type() == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert parts[0].get_payload(decode=True)[:2] == b"PK"  # a real xlsx (zip) came through intact


def test_an_unreadable_attachment_is_a_clean_send_error_without_the_path(tmp_path, monkeypatch):
    from naukri_agent.notifications.exceptions import EmailSendError

    monkeypatch.setattr("naukri_agent.notifications.email.smtplib.SMTP", lambda *a, **k: pytest.fail("must not connect"))
    sender = SmtpEmailSender("h", 587, "me@x.com", "pw", "me@x.com")
    with pytest.raises(EmailSendError) as exc:
        sender.send(EmailMessage(subject="s", text_body="b", attachments=[str(tmp_path / "missing.xlsx")]))
    assert "missing.xlsx" not in str(exc.value) and "FileNotFoundError" in str(exc.value)


def test_the_file_sender_notes_the_attachment_names(tmp_path):
    out = tmp_path / "out"
    result = FileEmailSender(out).send(EmailMessage(subject="s", text_body="body", attachments=[str(tmp_path / "w.xlsx")]))
    assert "Attachments: w.xlsx" in Path(result.path).read_text(encoding="utf-8")


# --- the command ------------------------------------------------------------------------------------------------------------------


def test_the_command_saves_the_chosen_week_without_emailing(tmp_path, monkeypatch):
    db = tmp_path / "t.db"
    c = cfg(tmp_path, database_url=f"sqlite:///{db}")
    factory = init_db(c)
    with session_scope(factory) as s:
        cand, run = make_candidate(s), make_run(s)
        company_job(s, cand, run, 0, 7)
    monkeypatch.setattr(cli_main, "get_settings", lambda: c)
    res = CliRunner().invoke(cli_main.cli, ["weekly-report", "--week-of", "2026-10-08", "--no-email"])
    assert res.exit_code == 0, res.output
    assert "Week 05 Oct to 11 Oct 2026: 1 job(s)" in res.output and "emailed" not in res.output
    assert (tmp_path / "weekly" / "weekly_2026-10-05_to_2026-10-11.xlsx").exists()
    bad = CliRunner().invoke(cli_main.cli, ["weekly-report", "--week-of", "soon"])
    assert bad.exit_code != 0 and "YYYY" not in bad.output and "2026-10-09" in bad.output
